from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin = s.bottom_margin = Inches(.7)
s.left_margin = s.right_margin = Inches(.8)
for name, size in [('Normal',12),('Title',24),('Subtitle',12),('Heading 1',18),('Heading 2',14),('List Bullet',12)]:
    st = doc.styles[name]
    st.font.name = 'Noto Sans Thai'
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0,0,0)
    rp=st.element.get_or_add_rPr()
    for attr in list(rp.rFonts.attrib): del rp.rFonts.attrib[attr]
    for attr in ['ascii','hAnsi','eastAsia','cs']: rp.rFonts.set(qn('w:'+attr),'Noto Sans Thai')
    sz=OxmlElement('w:szCs'); sz.set(qn('w:val'),str(size*2)); rp.append(sz)
    st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    for border in st.element.findall('.//'+qn('w:pBdr')): border.getparent().remove(border)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = Pt(17)
    if name in ['Heading 1','Heading 2']:
        st.paragraph_format.space_before=Pt(10)
        st.paragraph_format.line_spacing=1.05
doc.styles['Normal'].paragraph_format.widow_control = True
doc.styles['Title'].paragraph_format.line_spacing=1.15
h = s.header.paragraphs[0]
h.text = 'REMOTE STUDIO     คู่มือการใช้งานภาษาไทย'
h.runs[0].font.size = Pt(9)
f = s.footer.paragraphs[0]
f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
f.add_run('Remote Studio 0.1.0  •  ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); f._p.append(field)

def p(text,style=None): return doc.add_paragraph(text,style)
def title(text): doc.add_heading(text,1)
def sub(text): doc.add_heading(text,2)
def step(n,text): p(f'{n}. {text}')
def page(): doc.add_page_break()
def code(text):
    para=p(text)
    para.paragraph_format.line_spacing=1.05
    for r in para.runs:
        r.font.name='Menlo'; r.font.size=Pt(9)
def link(label,url):
    para=p('')
    el=OxmlElement('w:hyperlink')
    el.set(qn('r:id'),para.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
    run=OxmlElement('w:r'); props=OxmlElement('w:rPr'); color=OxmlElement('w:color'); color.set(qn('w:val'),'145A86'); props.append(color); run.append(props)
    t=OxmlElement('w:t'); t.text=label; run.append(t); el.append(run); para._p.append(el)

p('คู่มือการใช้งาน Remote Studio', 'Title')
p('การเริ่มต้นใช้งานและการเชื่อมกับ ChatGPT บนเว็บ','Subtitle')
p('สำหรับเจ้าของเครื่อง macOS Apple Silicon • รุ่น 0.1.0 Development preview\nตรวจสอบข้อมูลวันที่ 11 ตุลาคม 2026')
sub('สิ่งที่ควรรู้ก่อนเริ่ม')
p('แอปนี้เปิดให้ AI ใช้งานไฟล์และคำสั่งบนเครื่องผ่าน MCP โดยคุณเลือกโฟลเดอร์และตรวจอนุมัติการเปลี่ยนแปลงเอง หน้า Connections มี Set up ChatGPT Web สำหรับบันทึก config ตรวจ bridge และเปิด Secure MCP Tunnel')
p('กรอก Tunnel ID และ runtime API key ในแอป Key จะเก็บแบบเข้ารหัสด้วย OS และไม่ส่งกลับหน้า UI จากนั้นเพิ่ม plugin ในบัญชี ChatGPT ตามหน้า 5 การทดสอบในเครื่องผ่านแล้ว แต่ยังไม่ได้ยืนยันการเชื่อมกับบัญชีจริงของคุณ [1]')
sub('เลือกวิธีใช้ให้ตรงกับงาน')
p('ใช้ทันทีบนเครื่อง — เชื่อม local client แล้วให้ AI อ่านไฟล์ เสนอแก้ไข หรือรันงานผ่าน Remote Studio ตามหน้า 2 และ 3')
p('ใช้ ChatGPT เว็บแบบส่งข้อมูลเอง — วางข้อความหรือแนบไฟล์ที่ต้องการ แล้วนำคำตอบกลับมาแก้ไขบนเครื่อง เหมาะเมื่อยังไม่ตั้งค่า MCP ตามหน้า 4')
p('ให้ ChatGPT เว็บเรียกเครื่องมือจริง — ตั้งค่าในแอปตามหน้า 5 แล้วเพิ่ม plugin บนเว็บ ต้องมี tunnel และสิทธิ์ในบัญชีที่ใช้')
p('สร้างการเชื่อมผ่าน URL สำหรับใช้งานระยะยาว — ต้องพัฒนา HTTP MCP gateway และระบบยืนยันตัวตนเพิ่มเติม ตามหน้า 6')
sub('ขอบเขตของรุ่นนี้')
p('มีเครื่องมือจริง 21 รายการด้านไฟล์ คำสั่ง งานเบื้องหลัง โปรเจกต์ Git แบบอ่านอย่างเดียว และข้อมูลระบบ ส่วน Video Browser และ Documents ยังไม่เปิดใช้งาน ตัวติดตั้งที่มีคือ macOS ARM64 แบบยังไม่ลงนามและยังไม่ notarize')

page(); title('เริ่มใช้งานแอปบนเครื่อง')
sub('เปิดแอปและเพิ่มโฟลเดอร์')
step(1,'เปิด Remote Studio.app ที่สร้างไว้ หรือเปิด DMG แล้วคัดลอกแอปไป Applications ก่อนสร้าง connection เพื่อให้ตำแหน่งแอปคงที่ หาก macOS บล็อก ให้ตรวจสอบว่าเป็นไฟล์ที่คุณสร้างเองและใช้ตัวเลือกเปิดที่ระบบอนุญาต')
step(2,'ทำ onboarding แล้วไป Workspaces กด Add folder เลือกโฟลเดอร์ทดสอบ เช่น Documents/AIWork เริ่มจากไฟล์สำเนาที่ไม่มีข้อมูลลับ')
step(3,'โฟลเดอร์ใหม่เริ่มด้วย Read only ตรวจชื่อและ path ให้ถูกต้อง เปิดสิทธิ์เขียนเฉพาะเมื่อจะให้ AI เสนอแก้ไฟล์ และเปิดสิทธิ์ execute แยกเมื่อจะรันคำสั่ง')
step(4,'เปิดสวิตช์การเข้าถึงของ workspace ไว้ หากต้องการหยุดใช้งาน ให้ปิดหรือถอนสิทธิ์ในแอป')
sub('เชื่อมกับ local client')
step(1,'ไป Connections เลือก client ที่ต้องการ แล้วเลือกเฉพาะโฟลเดอร์ที่ client นี้ควรเข้าถึง')
step(2,'หากจะเขียนไฟล์ ให้เปิด Allow file change requests แล้วกด Create connection การเขียนยังต้องมีสิทธิ์ของ workspace และอนุมัติใน Activity')
step(3,'กด Test bridge เพื่อทดสอบ MCP จริง ถ้าขึ้น Handshake verified หมายถึงการทดสอบในแอปผ่าน ไม่ได้ยืนยันว่าโปรแกรม AI ภายนอกกำลังเชื่อมอยู่')
step(4,'สำหรับ Claude หรือ Cursor ใช้ Install configuration ได้ โดยแอปสำรองไฟล์เดิมและเก็บรายการอื่นไว้ หากมี remote-studio อยู่แล้ว จะไม่เขียนทับอัตโนมัติ สำหรับ Codex ใช้ TOML ที่แอปสร้าง')
step(5,'เปิดใหม่หรือ reconnect client แล้วขอให้เรียก workspace_list ตรวจว่ารายการตรงกับโฟลเดอร์ที่ให้สิทธิ์')
sub('สถานะที่อ่านได้จากแอป')
p('Configured คือสร้าง credential แล้ว • Handshake verified คือ diagnostic ผ่าน • Revoked คือถอนสิทธิ์แล้ว การแสดง tool call ล่าสุดอาจเกิดจาก diagnostic จึงควรยืนยันด้วยคำสั่งจาก client ที่ใช้งานจริงอีกครั้ง')
p('การสร้าง config ไม่ได้ติดตั้งโปรแกรม AI ให้ และแอปไม่ได้จำลองการเชื่อมต่อ ถ้าย้าย Remote Studio.app หลังตั้งค่าแล้ว ให้สร้างหรือแก้ config ให้ตรงกับตำแหน่งใหม่')

page(); title('อ่านไฟล์ แก้ไข และรันงาน')
sub('ให้ AI อ่านและเสนอแก้ไฟล์')
p('ตัวอย่าง prompt ใน local client ที่เชื่อมแล้ว')
p('“ใช้ Remote Studio เรียก workspace_list แล้วอ่าน README.md ในโฟลเดอร์ AIWork สรุปปัญหาให้ก่อน หากจะแก้ไข ให้เสนอการเปลี่ยนแปลงและรอฉันอนุมัติใน Activity”')
step(1,'AI ใช้ workspace ID จากผลลัพธ์จริง แล้วเรียก file_read หรือ file_search ภายในโฟลเดอร์ที่ได้รับสิทธิ์')
step(2,'เมื่อเรียก file_write หรือ file_patch ระบบสร้างข้อเสนอที่รออนุมัติ การได้รับ proposal ID ยังไม่ได้แปลว่าไฟล์เปลี่ยนแล้ว')
step(3,'ไป Activity อ่าน diff ตรวจ path และเนื้อหา แล้วกด Approve change เพื่อใช้การแก้ไข หรือปฏิเสธข้อเสนอ')
step(4,'ตรวจไฟล์ที่เปลี่ยนและผลลัพธ์จริง หากต้องการย้อนกลับ ใช้ Undo ระบบจะไม่เขียนทับไฟล์ที่มีการแก้เพิ่มหลังจากนั้น')
sub('ให้ AI รันคำสั่ง')
p('“ตรวจโปรเจกต์ใน AIWork แล้วเสนอรัน npm run build บอก executable arguments และโฟลเดอร์ทำงานก่อน รอฉันกดอนุมัติ และรายงาน exit code กับ log จริง”')
p('ต้องเปิด execute ใน workspace ทุกคำสั่งรอการอนุมัติที่ Activity ด้วย Approve & run จากนั้นตรวจสถานะและ log ของ job และยกเลิกได้ คำสั่งใช้ executable กับ arguments โดยไม่ผ่าน shell อัตโนมัติ')
p('สิทธิ์รันคำสั่งใช้สิทธิ์ของบัญชีระบบปฏิบัติการ cwd ไม่ใช่ sandbox คำสั่งอาจเข้าถึงตำแหน่งนอก workspace ได้ จึงต้องอ่านคำสั่งก่อนอนุมัติ')
sub('เมื่อปิดหน้าต่างหรือขาดการเชื่อมต่อ')
p('ปิดหน้าต่างแล้วแอปอยู่ใน tray งานยังทำต่อได้ แต่เลือก Quit จะหยุด agent การเชื่อมจาก AI ที่หลุดไม่ได้แปลว่างานสำเร็จ ให้ตรวจ job_status หรือ Activity หลังเชื่อมใหม่ หาก agent crash งานที่ค้างจะถูกระบุ interrupted และไม่รันซ้ำเอง')
p('เครื่องมือไฟล์ไม่อ่าน symlink และป้องกันตำแหน่งลับบางประเภท เช่น .env .ssh และข้อมูลภายในแอป อ่านไฟล์ข้อความได้ไม่เกิน 2 MiB อย่าแก้ข้อจำกัดด้วยการให้สิทธิ์ทั้งโฟลเดอร์บ้านโดยไม่จำเป็น')

page(); title('ใช้ ChatGPT บนเว็บโดยส่งข้อมูลเอง')
p('วิธีนี้เริ่มใช้ได้โดยไม่ตั้งค่า tunnel แต่ไม่ได้เชื่อม MCP กับ Remote Studio และ ChatGPT จะไม่ได้รับสิทธิ์อ่านโฟลเดอร์บนเครื่องเพียงเพราะคุณบอก path')
step(1,'เปิด ChatGPT บนเว็บ สร้างแชตใหม่ เลือกข้อความหรือไฟล์ที่คุณต้องการให้ช่วยวิเคราะห์ ถ้าบัญชีมีการแนบไฟล์ ให้แนบเฉพาะไฟล์นั้น หรือวางข้อความลงแชต')
step(2,'บอกงานและข้อกำหนดอย่างชัดเจน เช่น ต้องการสรุป ปรับ README หรือหาเหตุผลที่ build ล้มเหลว ส่ง log ที่เกี่ยวข้องและลบ secret ก่อนส่ง')
step(3,'ขอคำตอบเป็นเนื้อหาไฟล์เต็มหรือรายการแก้ไข ระบุว่า ChatGPT ต้องแยกข้อเสนอออกจากสิ่งที่ได้ทำจริง')
step(4,'ตรวจคำตอบ แล้วแก้ไฟล์ด้วย editor บนเครื่อง หรือใช้หน้าอ่านและแก้ไฟล์ใน Workspaces ของ Remote Studio ซึ่งจะสร้างข้อเสนอให้ตรวจใน Activity')
step(5,'หากต้องรันคำสั่ง ให้ตรวจแล้วใช้ Terminal บนเครื่อง หรือ local client ที่เชื่อม Remote Studio อย่าถือว่าข้อความ “รันเสร็จแล้ว” เป็นผลจากเครื่องของคุณโดยไม่มี tool result')
sub('ตัวอย่าง prompt ที่คัดลอกไปใช้ได้')
p('“นี่คือ README.md ของโปรเจกต์ ช่วยปรับวิธีติดตั้งให้ชัดเจน ส่งเนื้อหาไฟล์ฉบับเต็มกลับมา ฉันจะตรวจและบันทึกบนเครื่องเอง อย่าอ้างว่าได้แก้ไฟล์หรือรันคำสั่งบนเครื่องฉันแล้ว”')
p('“นี่คือ log จาก npm run build ช่วยระบุสาเหตุและเสนอขั้นตอนแก้ทีละข้อ พร้อมอธิบายผลที่ควรเห็นหลังแต่ละขั้นตอน”')
sub('สิ่งที่ไม่ได้เกิดขึ้นจากการวาง config')
p('JSON ของ Claude หรือ TOML ของ Codex เป็น config สำหรับ local client ไม่ใช่ URL ที่นำไปเพิ่มใน ChatGPT เว็บ การวาง credential หรือ path ของ socket ลงแชตไม่ได้สร้างการเชื่อมต่อ และไม่ควรส่ง credential ให้โมเดล')
sub('ถ้าต้องการให้เว็บเรียกเครื่องมือจริง')
p('เปิด Connections แล้วกด Set up ChatGPT Web ตามหน้าถัดไป กรอกค่าในแอปแทนการส่ง key ลงแชต แอปจัดการ runtime และ health checks ให้ ส่วนการเพิ่ม plugin ต้องทำในบัญชี ChatGPT ของคุณ')

page(); title('ตั้งค่า ChatGPT Web ในแอป')
p('ใช้ OpenAI Secure MCP Tunnel กับ stdio bridge โดยแอปจัดการ command และ local credential ให้ กรอก key เฉพาะในแอปและเพิ่ม plugin ตามขั้นตอนบนเว็บ [1]')
step(1,'เพิ่มโฟลเดอร์ใน Workspaces จากนั้นเปิด Connections กด Set up ChatGPT Web')
step(2,'กรอก Tunnel ID ด้วยค่า LOCALBRIDGE_TUNNEL_ID และ Runtime API key ด้วยค่า CONTROL_PLANE_API_KEY ใช้ runtime key ที่มี Tunnels Read + Use Tunnel ต้องผูกกับ ChatGPT workspace ที่จะใช้งาน')
step(3,'ตรวจ Official tunnel-client แอปจะตรวจพบ Homebrew บน macOS หากยังไม่มี ให้ติดตั้งด้วยคำสั่งด้านล่าง แล้วกด Choose executable เลือก /opt/homebrew/bin/tunnel-client หรือ /usr/local/bin/tunnel-client')
code('brew install openai/tools/tunnel-client')
step(4,'เลือก Allowed folders เฉพาะที่ต้องการเริ่มทดลอง ไม่มีโฟลเดอร์ที่เลือกคือไม่มีสิทธิ์อ่านไฟล์ Allow file change requests เปิดภายหลังได้และยังต้องอนุมัติใน Activity ข้อมูลที่เครื่องมืออ่านจะส่งเป็นผลลัพธ์ไปยัง ChatGPT')
step(5,'กด Save & start แอปเข้ารหัส key ตรวจ MCP bridge และเปิด tunnel หากต้องการเปิดอัตโนมัติทุกครั้งที่เปิดแอป ให้เปิด Start tunnel when Remote Studio opens')
step(6,'ตรวจ Local bridge และ Tunnel runtime แล้วดู OpenAI polling Runtime ready ยืนยัน health ของเครื่องนี้ ส่วนเวลาตอบกลับจาก OpenAI แสดงแยกกัน ไม่ได้ยืนยันว่าติดตั้ง plugin บนเว็บแล้ว')
step(7,'กด Open ChatGPT ไป Plugins กด + แล้ว Add custom MCP server เลือก Connection เป็น Tunnel เลือกหรือวาง Tunnel ID ตั้งชื่อ Remote Studio ตั้ง authentication ตามเซิร์ฟเวอร์ ตรวจคำเตือนแล้ว Create as a plugin [2]')
p('Bridge นี้ไม่มี OAuth เพิ่มเติมและอาศัยสิทธิ์ของ tunnel กับ workspace เปิดแอปไว้หรือปิดหน้าต่างให้ทำงานต่อใน menu bar ปุ่ม Stop tunnel หยุดการรับคำขอ และ Quit จะหยุดทั้ง tunnel และ agent')

page(); title('ตรวจการเชื่อมต่อและทางเลือกเพิ่มเติม')
sub('ทดสอบจาก ChatGPT ก่อนให้เขียน')
p('ติดตั้ง plugin ที่สร้าง เปิดแชตใหม่ พิมพ์ @ เลือก plugin แล้วขอให้เรียก workspace_list ตรวจ tool result และโฟลเดอร์ที่ได้จริง [3] อ่านไฟล์ทดสอบหนึ่งไฟล์ก่อน เมื่อผ่านแล้วค่อยเพิ่มสิทธิ์เขียนและตรวจ approval ใน Activity ทุกครั้ง')
p('ตัวอย่าง “ใช้ Remote Studio เรียก workspace_list แล้วอ่าน README.md ใน workspace ทดสอบ รายงานชื่อไฟล์และผลลัพธ์จริง ยังไม่แก้ไขและไม่รันคำสั่ง”')
p('การอนุมัติของ ChatGPT และการอนุมัติใน Remote Studio เป็นคนละขั้นตอน ผลที่ระบุ pending หรือ proposal ต้องกลับไปอนุมัติบนเครื่อง ก่อนให้ AI ตรวจผลอีกครั้ง')
sub('แก้ปัญหาที่พบบ่อย')
p('ไม่มี Tunnel หรือเพิ่ม custom MCP ไม่ได้ — ตรวจสิทธิ์ workspace นโยบายความปลอดภัย และการผูก tunnel กับ workspace อย่าคาดว่าทุกบัญชีเปิดได้เหมือนกัน')
p('ค้นหา tools ไม่ได้ — เปิดแอป ตรวจว่าไม่ได้ Quit กด Start tunnel และ Run checks ดู Runtime diagnostics ในหน้า Connections หาก OpenAI polling ไม่มีผลสำเร็จ ให้ตรวจ key สิทธิ์ Read + Use และ Tunnel ID')
p('อ่านไม่ได้หรือรายการว่าง — ตรวจ workspace enabled scope ของ connection และ credential ที่ยังไม่ revoked หาก path เป็น symlink หรือไฟล์ที่ระบบป้องกัน ให้ใช้ไฟล์ทดสอบปกติ')
p('แก้ไขไม่ได้ — ตรวจสิทธิ์เขียนทั้ง workspace และ client จากนั้นไป Activity หากไฟล์เปลี่ยนหลังเสนอแก้ ให้สร้างข้อเสนอใหม่จากเนื้อหาปัจจุบัน')
p('รายการ tools เก่า — หยุดและเริ่ม tunnel ใหม่ แล้ว refresh plugin หรือเปิดแชตใหม่หลังเปลี่ยนโมดูล [3] หากเลิกใช้ ให้ Edit ChatGPT settings แล้ว Remove เพื่อลบ key ในเครื่องและ revoke local grant')
sub('หากต้องการเชื่อมผ่าน HTTPS URL')
p('ต้องเพิ่ม HTTP MCP gateway ที่มี authentication และส่งคำขอผ่าน scoped credential ไปยัง agent เดิม โดยคง local approval และการตรวจสิทธิ์ไว้ จากนั้นใช้ endpoint ที่เข้าถึงได้และเพิ่ม Server URL ใน ChatGPT [2] วิธีนี้ยังไม่ได้พัฒนาในรุ่น 0.1.0 ไม่ควรเปิด socket ของ agent หรือ UI admin credential สู่สาธารณะ')
sub('แหล่งอ้างอิงและขอบเขตการตรวจสอบ')
link('[1] OpenAI Secure MCP Tunnel','https://developers.openai.com/api/docs/guides/secure-mcp-tunnels')
link('[2] OpenAI Add custom MCP server','https://developers.openai.com/api/docs/guides/custom-mcp-server')
link('[3] OpenAI Connect and test your plugin','https://developers.openai.com/plugins/deploy/connect-chatgpt')

doc.core_properties.title='คู่มือการใช้งาน Remote Studio และ ChatGPT บนเว็บ'
doc.core_properties.subject='Remote Studio 0.1.0 Thai user manual'
out=ROOT/'docs'/'Remote-Studio-User-Guide-TH.docx'
doc.save(out)
print(out)
